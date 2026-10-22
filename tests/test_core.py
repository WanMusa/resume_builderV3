import copy
import json
import os
import shutil
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from urllib.parse import unquote

import pytest
import httpx
from docx import Document
from pydantic import ValidationError
from sqlalchemy import func, select, update
from sqlalchemy.dialects import postgresql
from sqlalchemy.engine import URL
from sqlalchemy.schema import CreateSchema, DropSchema

from app.config import Settings
from app.contracts import EnvelopeContract, JobContract
from app.db import init_db, make_engine, make_session_factory
from app.importer import EnvelopeValidationError, import_envelope, scan_inbox
from app.models import Artifact, Batch, GeneratedDocument, Job, JobInput, utcnow
from app.worker import (LostLease, _fence, claim_job, process_one, regenerate_job,
                        rerender_job, retry_job)


@pytest.fixture
def workspace(tmp_path):
    return tmp_path


@pytest.fixture
def database(workspace):
    settings = Settings(
        _env_file=None, database_url=os.environ.get("CORE_TEST_DATABASE_URL", f"sqlite:///{workspace / 'core.db'}"),
        database_password=os.environ.get("CORE_TEST_DATABASE_PASSWORD"),
        supabase_bucket="test-documents",
        data_dir=workspace / "data", inbox_dir=workspace / "inbox", max_attempts=2,
    )
    engine = make_engine(settings)
    schema = None
    if engine.dialect.name == "postgresql":
        schema = f"core_test_{uuid4().hex}"
        with engine.begin() as connection:
            connection.execute(CreateSchema(schema))
        engine = engine.execution_options(schema_translate_map={None: schema})
    init_db(engine)
    yield make_session_factory(engine), settings
    if schema:
        with engine.begin() as connection:
            connection.execute(DropSchema(schema, cascade=True))
    engine.dispose()


@pytest.fixture
def payload():
    return {"schema_version": 2, "jobs": [{
        "id": "job-1", "source_url": "https://example.com/jobs/1", "site": "example",
        "captured_at": "2026-10-14T12:00:00Z", "description_text": "Build reliable software.",
    }]}


def intake(factory, payload):
    with factory() as session:
        batch = import_envelope(session, payload, "test")
        session.commit()
        return batch.id


def get_job(factory):
    with factory() as session:
        return session.scalar(select(Job))


@pytest.fixture
def pipeline(monkeypatch):
    calls = {"generation": 0, "render": 0, "upload": 0}
    package = {"resume": {}, "cover_letter": {}, "provenance": {"provider": "mock"}}

    def generate(payload, settings):
        calls["generation"] += 1
        return copy.deepcopy(package)

    def render(package, assets_dir, job):
        calls["render"] += 1
        return {"resume.docx": b"Test resume"}

    class Storage:
        def check_bucket(self):
            pass

        def upload(self, content, key):
            calls["upload"] += 1
            return {"backend": "supabase", "bucket": settings_bucket,
                    "object_key": key, "local_path": None}

    settings_bucket = "test-documents"
    monkeypatch.setattr("app.generation.generate_package", generate)
    monkeypatch.setattr("app.rendering.render_package", render)
    monkeypatch.setattr("app.storage.get_storage", lambda settings: Storage())
    return calls, package


def test_defaults_and_postgres_models():
    settings = Settings(_env_file=None)
    assert settings.database_url == "postgresql+psycopg://resume_builder@localhost:5432/resume_builder"
    assert settings.provider == "mock"
    assert settings.database_password is None
    assert not settings.allow_unauthenticated
    assert settings.lease_seconds == 600
    assert Job.__table__.c.payload.type.compile(dialect=postgresql.dialect()) == "JSONB"
    with pytest.raises(ValidationError):
        Settings(_env_file=None, max_attempts=0)
    with pytest.raises(ValidationError):
        Settings(_env_file=None, lease_seconds=299)


def test_dotenv_password_passed_to_postgres_driver(workspace, monkeypatch):
    from sqlalchemy import create_engine as original_create_engine
    from app import db

    dotenv = workspace / "settings.env"
    dotenv.write_text("DATABASE_PASSWORD=fixture-value\n")
    settings = Settings(_env_file=dotenv)
    captured = {}

    def capture(url, **kwargs):
        captured.update(kwargs)
        return original_create_engine(url, **kwargs)

    monkeypatch.setattr(db, "create_engine", capture)
    engine = make_engine(settings)
    assert settings.database_password == "fixture-value"
    assert captured["connect_args"] == {"password": "fixture-value"}
    assert engine.url.password is None
    engine.dispose()


def test_database_password_does_not_affect_sqlite(workspace):
    settings = Settings(
        _env_file=None, database_url=f"sqlite:///{workspace / 'password.db'}",
        database_password="fixture-value",
    )
    engine = make_engine(settings)
    init_db(engine)
    engine.dispose()


def test_custom_inline_password_url_preserved_without_separate_password():
    fixture_value = "fixture-value"
    url = URL.create(
        "postgresql+psycopg", "example", fixture_value, "localhost", None, "example",
    ).render_as_string(hide_password=False)
    engine = make_engine(Settings(_env_file=None, database_url=url, database_password=None))
    assert engine.url.password == "fixture-value"
    engine.dispose()


def test_strict_contract_defaults(payload):
    job = JobContract.model_validate(payload["jobs"][0])
    assert job.title is None and job.notes is None and job.warnings == []
    assert job.requirements == [] and job.responsibilities == [] and job.skills == []
    assert job.extraction_mode == "raw"
    assert EnvelopeContract.model_validate(payload).exported_at is None


@pytest.mark.parametrize(("field", "value"), [
    ("id", 42), ("site", False), ("description_text", None), ("captured_at", 42),
    ("source_url", "file:///etc/passwd"), ("source_url", "https:///missing-host"),
    ("captured_at", "2026-10-14"), ("title", 42), ("notes", ["ok", 12]),
    ("extraction_mode", "other"), ("requirements", ["ok", 42]),
    ("responsibilities", "wrong"), ("skills", [False]),
])
def test_invalid_job_types(payload, field, value):
    payload["jobs"][0][field] = value
    with pytest.raises(ValidationError):
        JobContract.model_validate(payload["jobs"][0])


def test_refined_contract_and_nullable_notes(payload):
    job = payload["jobs"][0]
    job.update(extraction_mode="llm_refined", notes="Review extraction",
               requirements=["Five years experience"], responsibilities=["Build services"],
               skills=["Python"], warnings=["Extraction needs review"])
    validated = JobContract.model_validate(job)
    assert validated.notes == "Review extraction"
    assert validated.extraction_mode == "llm_refined"
    assert validated.skills == ["Python"]
    job["notes"] = None
    assert JobContract.model_validate(job).notes is None


@pytest.mark.parametrize("invalid", [
    {"schema_version": 1, "jobs": [{}]},
    {"schema_version": "2", "jobs": [{}]},
    {"schema_version": 2.0, "jobs": [{}]},
    {"schema_version": 2, "jobs": []},
    {"schema_version": 2, "jobs": "wrong"},
    {"schema_version": 2, "jobs": [{}], "unexpected": True},
])
def test_invalid_envelope(database, invalid):
    factory, _ = database
    with factory() as session, pytest.raises(EnvelopeValidationError) as exc:
        import_envelope(session, invalid, "test")
    assert exc.value.errors


def test_partial_import_and_all_invalid(database, payload):
    factory, _ = database
    payload["jobs"].append({"id": "bad", "description_text": 42})
    batch_id = intake(factory, payload)
    with factory() as session:
        batch = session.get(Batch, batch_id)
        assert batch.status == "partial"
        assert batch.errors[0]["index"] == 1
        assert len(batch.jobs) == 1
        assert batch.jobs[0].payload["title"] is None
        assert batch.jobs[0].input.payload == batch.jobs[0].payload
        with pytest.raises(EnvelopeValidationError):
            import_envelope(session, {"schema_version": 2, "jobs": [{}]}, "test")


def test_validation_errors_do_not_persist_submitted_input(database, payload):
    factory, settings = database
    secret_input = {"private": "personal-secret-value"}
    payload["jobs"].append({**payload["jobs"][0], "description_text": secret_input})
    batch_id = intake(factory, payload)
    with factory() as session:
        errors = session.get(Batch, batch_id).errors
        assert errors[0]["errors"][0]["loc"] == ["description_text"]
        assert "personal-secret-value" not in json.dumps(errors)
        assert "input" not in errors[0]["errors"][0]
        with pytest.raises(EnvelopeValidationError) as exc:
            import_envelope(session, {"schema_version": 2, "jobs": secret_input}, "test")
        assert "personal-secret-value" not in str(exc.value)
    settings.inbox_dir.mkdir()
    (settings.inbox_dir / "invalid.json").write_text(json.dumps({
        "schema_version": 2, "jobs": [{"description_text": secret_input}],
    }))
    scan_inbox(factory, settings)
    errors_path = settings.inbox_dir / "quarantine" / "invalid.errors.json"
    assert "personal-secret-value" not in errors_path.read_text()


def test_deduplication_and_changed_revision(database, payload):
    factory, _ = database
    first = intake(factory, payload)
    assert intake(factory, payload) == first
    revised = copy.deepcopy(payload)
    revised["jobs"][0]["captured_at"] = "2026-10-15T12:00:00Z"
    revised["jobs"][0]["description_text"] = " Build   reliable software.\n"
    second = intake(factory, revised)
    assert second != first
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(Job)) == 1
        assert session.get(Batch, first).jobs[0].id == session.get(Batch, second).jobs[0].id
        snapshot = session.scalar(select(JobInput))
        assert snapshot.payload["captured_at"] == payload["jobs"][0]["captured_at"]
    revised["jobs"][0]["description_text"] = "Build entirely different software."
    intake(factory, revised)
    with factory() as session:
        jobs = list(session.scalars(select(Job)))
        assert len(jobs) == 2 and jobs[0].id != jobs[1].id


def test_concurrent_import(database, payload):
    factory, _ = database
    barrier = threading.Barrier(4)

    def run(_):
        barrier.wait()
        return intake(factory, payload)

    with ThreadPoolExecutor(max_workers=4) as pool:
        ids = list(pool.map(run, range(4)))
    assert len(set(ids)) == 1
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(Job)) == 1
        assert session.scalar(select(func.count()).select_from(Batch)) == 1
        assert session.scalar(select(func.count()).select_from(JobInput)) == 1


def test_scan_archives_partial_and_quarantines_invalid(database, payload):
    factory, settings = database
    settings.inbox_dir.mkdir()
    payload["jobs"].append({"id": 123})
    (settings.inbox_dir / "valid.json").write_text(json.dumps(payload))
    (settings.inbox_dir / "bad.json").write_text("{invalid")
    (settings.inbox_dir / "upload.part").write_text("still copying")
    results = scan_inbox(factory, settings)
    assert {r["status"] for r in results} == {"partial", "quarantined"}
    assert (settings.inbox_dir / "archive" / "valid.json").exists()
    assert (settings.inbox_dir / "quarantine" / "bad.json").exists()
    assert json.loads((settings.inbox_dir / "quarantine" / "bad.errors.json").read_text())
    assert (settings.inbox_dir / "upload.part").exists()
    assert scan_inbox(factory, settings) == []


def test_worker_success(database, payload, pipeline):
    factory, settings = database
    calls, _ = pipeline
    intake(factory, payload)
    assert process_one(factory, settings)
    assert not process_one(factory, settings)
    job = get_job(factory)
    assert job.status == "completed" and job.attempts == 1
    assert job.lease_token is None and job.completed_at is not None
    with factory() as session:
        assert session.scalar(select(GeneratedDocument)).package["resume"] == {}
        artifact = session.scalar(select(Artifact))
        assert artifact.object_key == "resume.docx"
        assert artifact.backend == "supabase" and artifact.local_path is None
    assert not settings.data_dir.exists()
    assert calls == {"generation": 1, "render": 1, "upload": 1}


def test_review_required(database, payload, pipeline):
    factory, settings = database
    _, package = pipeline
    package["review_required"] = True
    intake(factory, payload)
    assert process_one(factory, settings)
    job = get_job(factory)
    assert job.status == "needs_review"
    with factory() as session, pytest.raises(ValueError, match="Only failed jobs"):
        retry_job(session, job.id)


def test_checkpoint_survives_render_failure(database, payload, pipeline, monkeypatch):
    factory, settings = database
    calls, _ = pipeline
    intake(factory, payload)
    from app.rendering import render_package

    def fail(package, directory, job):
        # Generation must already be committed, visible from another session.
        with factory() as session:
            assert session.scalar(select(GeneratedDocument)) is not None
        raise RuntimeError("renderer unavailable")

    monkeypatch.setattr("app.rendering.render_package", fail)
    assert process_one(factory, settings)
    assert get_job(factory).status == "queued"
    monkeypatch.setattr("app.rendering.render_package", render_package)
    assert process_one(factory, settings)
    assert get_job(factory).status == "completed"
    assert calls["generation"] == 1


def test_worker_error_does_not_persist_exception_response(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)

    def fail(payload, settings):
        raise ValueError("provider response includes private-personal-content and secret-key")

    monkeypatch.setattr("app.generation.generate_package", fail)
    assert process_one(factory, settings)
    job = get_job(factory)
    assert job.status == "queued"
    assert job.error == "ValueError: Job processing failed"
    assert "private-personal-content" not in job.error
    assert "secret-key" not in job.error


@pytest.mark.parametrize("reason", ["unknown_evidence", "content_length"])
def test_worker_reports_safe_evidence_validation(database, payload, pipeline, monkeypatch, reason):
    from app.generation import EvidenceValidationError

    factory, settings = database
    intake(factory, payload)

    def fail(payload, settings):
        raise EvidenceValidationError(reason)

    monkeypatch.setattr("app.generation.generate_package", fail)
    assert process_one(factory, settings)
    job = get_job(factory)
    assert job.status == "queued"
    assert job.error == f"Evidence validation [{reason}]: {EvidenceValidationError.MESSAGES[reason]}"
    with factory() as session:
        assert session.scalar(select(GeneratedDocument)) is None


def test_bounded_attempts_and_manual_retry(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)
    from app.rendering import render_package

    def fail(*args):
        raise RuntimeError("deterministic render error")

    monkeypatch.setattr("app.rendering.render_package", fail)
    process_one(factory, settings)
    process_one(factory, settings)
    job = get_job(factory)
    assert job.status == "failed" and job.attempts == 2
    assert job.error == "RuntimeError: Job processing failed"
    assert not process_one(factory, settings)
    with factory() as session:
        retry_job(session, job.id)
        session.commit()
    assert get_job(factory).attempts == 0
    monkeypatch.setattr("app.rendering.render_package", render_package)
    assert process_one(factory, settings)
    assert get_job(factory).status == "completed"
    assert pipeline[0]["generation"] == 1
    with factory() as session, pytest.raises(ValueError):
        retry_job(session, job.id)


def test_expired_claim_fences_stale_worker(database, payload):
    factory, settings = database
    intake(factory, payload)
    first = claim_job(factory, settings)
    assert claim_job(factory, settings) is None
    with factory() as session:
        session.execute(update(Job).values(lease_until=utcnow() - timedelta(seconds=1)))
        session.commit()
    second = claim_job(factory, settings)
    assert first["token"] != second["token"]
    with factory() as session, pytest.raises(LostLease):
        _fence(session, first["id"], first["token"], settings, status="completed")
    assert get_job(factory).lease_token == second["token"]


def test_concurrent_workers_claim_only_once(database, payload):
    factory, settings = database
    intake(factory, payload)
    barrier = threading.Barrier(4)

    def claim(_):
        barrier.wait()
        return claim_job(factory, settings)

    with ThreadPoolExecutor(max_workers=4) as pool:
        claims = list(pool.map(claim, range(4)))
    assert sum(claim is not None for claim in claims) == 1
    assert get_job(factory).attempts == 1


def test_postgres_worker_skips_locked_candidate(database, payload):
    factory, settings = database
    if factory.kw["bind"].dialect.name != "postgresql":
        pytest.skip("Requires CORE_TEST_DATABASE_URL pointing to PostgreSQL")
    payload["jobs"].append({**payload["jobs"][0], "id": "job-2"})
    intake(factory, payload)
    with factory() as session:
        locked = session.scalar(select(Job).order_by(Job.created_at, Job.id).limit(1).with_for_update())
        locked_id = locked.id
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(claim_job, factory, settings)
            try:
                claimed = future.result(timeout=5)
            finally:
                session.rollback()
    assert claimed is not None and claimed["id"] != locked_id
    with factory() as session:
        assert session.get(Job, locked_id).status == "queued"


def test_expired_renderer_cannot_upload(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)

    def stale_render(package, directory, job):
        with factory() as session:
            session.execute(update(Job).values(lease_until=utcnow() - timedelta(seconds=1)))
            session.commit()
        assert claim_job(factory, settings) is not None
        return {"resume.docx": b"stale artifact"}

    monkeypatch.setattr("app.rendering.render_package", stale_render)
    assert process_one(factory, settings)
    assert pipeline[0]["upload"] == 0
    with factory() as session:
        assert session.scalar(select(Artifact)) is None
        assert session.scalar(select(GeneratedDocument)) is not None


def test_stale_upload_cannot_persist_artifact_metadata(database, payload, pipeline, monkeypatch):
    factory, settings = database
    if factory.kw["bind"].dialect.name == "postgresql":
        pytest.skip("SQLite fault injection; PostgreSQL upload locking has separate coverage")
    intake(factory, payload)
    successor = {}

    class Storage:
        def check_bucket(self):
            pass

        def upload(self, content, key):
            with factory() as session:
                session.execute(update(Job).values(lease_until=utcnow() - timedelta(seconds=1)))
                session.commit()
            successor.update(claim_job(factory, settings))
            assert key == "resume.docx"
            return {"backend": "supabase", "object_key": key, "local_path": None}

    monkeypatch.setattr("app.storage.get_storage", lambda settings: Storage())
    assert process_one(factory, settings)
    with factory() as session:
        assert session.scalar(select(Artifact)) is None
        assert session.scalar(select(GeneratedDocument)) is not None
        assert session.scalar(select(Job)).lease_token == successor["token"]


def test_expired_last_attempt_is_terminal(database, payload):
    factory, settings = database
    settings.max_attempts = 1
    intake(factory, payload)
    claim_job(factory, settings)
    with factory() as session:
        session.execute(update(Job).values(lease_until=utcnow() - timedelta(seconds=1)))
        session.commit()
    assert claim_job(factory, settings) is None
    assert get_job(factory).status == "failed"


def test_stale_generation_cannot_checkpoint(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)

    def stale_generate(payload, settings):
        with factory() as session:
            session.execute(update(Job).values(lease_until=utcnow() - timedelta(seconds=1)))
            session.commit()
        replacement = claim_job(factory, settings)
        assert replacement is not None
        return pipeline[1]

    monkeypatch.setattr("app.generation.generate_package", stale_generate)
    assert process_one(factory, settings)
    with factory() as session:
        assert session.scalar(select(GeneratedDocument)) is None
    assert get_job(factory).attempts == 2


def test_upload_retry_keeps_checkpoint_and_completed_artifacts(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)
    upload_names = []

    def render(package, directory, job):
        return {"resume.docx": b"resume", "cover_letter.docx": b"cover letter"}

    class Storage:
        failed = False

        def check_bucket(self):
            pass

        def upload(self, content, key):
            name = key.rsplit("/", 1)[-1]
            upload_names.append(name)
            if name == "cover_letter.docx" and not self.failed:
                self.failed = True
                raise RuntimeError("storage unavailable")
            return {"backend": "supabase", "bucket": settings.supabase_bucket,
                    "object_key": key, "local_path": None}

    storage = Storage()
    monkeypatch.setattr("app.rendering.render_package", render)
    monkeypatch.setattr("app.storage.get_storage", lambda settings: storage)
    process_one(factory, settings)
    assert get_job(factory).status == "queued"
    process_one(factory, settings)
    assert get_job(factory).status == "completed"
    assert pipeline[0]["generation"] == 1
    assert upload_names == ["resume.docx", "cover_letter.docx", "cover_letter.docx"]
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(Artifact)) == 2


@pytest.mark.parametrize("status", ["completed", "needs_review", "failed"])
def test_rerender_uses_checkpoint_without_generation(database, payload, pipeline, status):
    factory, settings = database
    calls, _ = pipeline
    intake(factory, payload)
    process_one(factory, settings)
    job = get_job(factory)
    with factory() as session:
        checkpoint = session.scalar(select(GeneratedDocument))
        checkpoint_id, original = checkpoint.id, copy.deepcopy(checkpoint.package)
        session.execute(update(Job).values(status=status))
        session.commit()
        rerender_job(session, job.id)
        session.commit()
        assert session.scalar(select(Artifact)) is None
    assert process_one(factory, settings)
    assert calls["generation"] == 1
    assert calls["render"] == 2 and calls["upload"] == 2
    with factory() as session:
        checkpoint = session.scalar(select(GeneratedDocument))
        assert checkpoint.id == checkpoint_id and checkpoint.package == original
    assert not settings.data_dir.exists()


def test_rerender_requires_checkpoint_and_terminal_status(database, payload, pipeline):
    factory, settings = database
    intake(factory, payload)
    job = get_job(factory)
    with factory() as session:
        with pytest.raises(ValueError, match="checkpoint"):
            rerender_job(session, job.id)
        with pytest.raises(LookupError):
            rerender_job(session, "missing")
    process_one(factory, settings)
    with factory() as session:
        rerender_job(session, job.id)
        session.commit()
        with pytest.raises(ValueError, match="terminal"):
            rerender_job(session, job.id)


def test_bucket_failure_prevents_generation_and_local_output(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)

    class Storage:
        def check_bucket(self):
            raise RuntimeError("Storage bucket check returned HTTP 403")

    monkeypatch.setattr("app.storage.get_storage", lambda settings: Storage())
    process_one(factory, settings)
    assert pipeline[0] == {"generation": 0, "render": 0, "upload": 0}
    assert not settings.data_dir.exists()
    assert get_job(factory).status == "queued"


def test_retry_upload_replaces_other_bucket_metadata(database, payload, pipeline):
    factory, settings = database
    intake(factory, payload)
    job = get_job(factory)
    with factory() as session:
        session.add(GeneratedDocument(job_id=job.id, package=copy.deepcopy(pipeline[1])))
        session.add(Artifact(job_id=job.id, kind="resume", filename="resume.docx",
                             backend="supabase", bucket="old-bucket", object_key="old/resume.docx"))
        session.commit()
    process_one(factory, settings)
    assert get_job(factory).status == "completed"
    assert pipeline[0]["generation"] == 0 and pipeline[0]["upload"] == 1
    with factory() as session:
        artifacts = list(session.scalars(select(Artifact)))
        assert len(artifacts) == 1
        assert artifacts[0].bucket == settings.supabase_bucket
        assert artifacts[0].object_key != "old/resume.docx"


def test_checkpoint_to_word_to_supabase_without_llm_or_local_files(database, payload, monkeypatch):
    from app.generation import generate_package
    from app.storage import DOCX_CONTENT_TYPE

    factory, settings = database
    root = Path(__file__).resolve().parents[1]
    settings.assets_dir = settings.data_dir.parent / "assets"
    settings.assets_dir.mkdir()
    shutil.copyfile(root / "assets/base_resume_example.json", settings.assets_dir / "base_resume.json")
    shutil.copytree(root / "assets/templates", settings.assets_dir / "templates")
    settings.prompts_dir = root / "prompts"
    settings.supabase_url = "https://storage.example"
    settings.supabase_key = "test-only"
    package = generate_package(payload["jobs"][0], settings)
    package["review_required"] = True
    intake(factory, payload)
    job = get_job(factory)
    with factory() as session:
        session.get(Job, job.id).status = "needs_review"
        session.add(GeneratedDocument(job_id=job.id, package=package))
        session.commit()
        rerender_job(session, job.id)
        session.commit()

    def forbidden(*args, **kwargs):
        pytest.fail("Re-render must never invoke generation")

    uploaded = {}

    def upload(self, url, **kwargs):
        assert kwargs["headers"]["Content-Type"] == DOCX_CONTENT_TYPE
        name = unquote(url.rsplit("/", 1)[-1])
        uploaded[name] = Document(BytesIO(kwargs["content"]))
        return httpx.Response(200, json={})

    monkeypatch.setattr("app.generation.generate_package", forbidden)
    monkeypatch.setattr(httpx.Client, "get",
                        lambda *args, **kwargs: httpx.Response(200, json={"public": False}))
    monkeypatch.setattr(httpx.Client, "post", upload)
    assert process_one(factory, settings)
    assert get_job(factory).status == "needs_review"
    assert len(uploaded) == 2
    assert any(name.endswith("-Resume.docx") for name in uploaded)
    assert any(name.endswith("-Cover Letter.docx") for name in uploaded)
    assert not settings.data_dir.exists()
    assert not list(settings.assets_dir.parent.glob("*.docx"))
    with factory() as session:
        assert session.scalar(select(GeneratedDocument)).package == package
        artifacts = list(session.scalars(select(Artifact)))
        assert len(artifacts) == 2
        assert all(a.backend == "supabase" and a.local_path is None for a in artifacts)
        assert {a.kind for a in artifacts} == {"resume", "cover_letter"}
        assert all(a.object_key == a.filename and "/" not in a.object_key for a in artifacts)


def test_regeneration_preserves_old_versions_and_checkpoints_retry(database, payload, pipeline, monkeypatch):
    factory, settings = database
    calls, _ = pipeline
    intake(factory, payload)
    process_one(factory, settings)
    job = get_job(factory)
    with factory() as session:
        original = copy.deepcopy(session.scalar(select(GeneratedDocument)).package)
        regenerate_job(session, job.id)
        session.commit()
        checkpoint = session.scalar(select(GeneratedDocument))
        assert checkpoint.package == {**original, "regeneration_requested": True}
        assert session.scalar(select(Artifact)) is None
    original_render = __import__("app.rendering", fromlist=["render_package"]).render_package

    def fail(*args):
        raise RuntimeError("Render failure after paid generation")

    monkeypatch.setattr("app.rendering.render_package", fail)
    process_one(factory, settings)
    assert calls["generation"] == 2
    with factory() as session:
        checkpoint = session.scalar(select(GeneratedDocument))
        assert checkpoint.package["previous_versions"] == [original]
        assert not checkpoint.package.get("regeneration_requested")
    monkeypatch.setattr("app.rendering.render_package", original_render)
    process_one(factory, settings)
    assert get_job(factory).status == "completed"
    assert calls["generation"] == 2


def test_failed_regeneration_does_not_erase_saved_content(database, payload, pipeline, monkeypatch):
    factory, settings = database
    intake(factory, payload)
    process_one(factory, settings)
    job = get_job(factory)
    with factory() as session:
        original = copy.deepcopy(session.scalar(select(GeneratedDocument)).package)
        regenerate_job(session, job.id)
        session.commit()

    def fail(*args):
        raise RuntimeError("Provider unavailable")

    monkeypatch.setattr("app.generation.generate_package", fail)
    process_one(factory, settings)
    process_one(factory, settings)
    assert get_job(factory).status == "failed"
    with factory() as session:
        assert session.scalar(select(GeneratedDocument)).package == {
            **original, "regeneration_requested": True}
        with pytest.raises(ValueError, match="Regeneration is pending"):
            rerender_job(session, job.id)
        retry_job(session, job.id)
        session.commit()
    assert get_job(factory).status == "queued"


def test_legacy_render_failure_is_actionable_without_paid_calls(database, payload, pipeline, monkeypatch):
    from app.rendering import PackageFormatError

    factory, settings = database
    intake(factory, payload)
    with factory() as session:
        session.add(GeneratedDocument(job_id=get_job(factory).id,
                                      package=copy.deepcopy(pipeline[1])))
        session.commit()

    def fail(*args):
        raise PackageFormatError()

    monkeypatch.setattr("app.rendering.render_package", fail)
    process_one(factory, settings)
    assert pipeline[0]["generation"] == 0
    assert "/regenerate" in get_job(factory).error


def test_postgres_upload_lock_blocks_expired_lease_reclaim(database, payload, pipeline, monkeypatch):
    factory, settings = database
    if factory.kw["bind"].dialect.name != "postgresql":
        pytest.skip("Requires CORE_TEST_DATABASE_URL pointing to PostgreSQL")
    intake(factory, payload)
    later = utcnow() + timedelta(seconds=settings.lease_seconds + 60)

    class Storage:
        def check_bucket(self):
            pass

        def upload(self, content, key):
            assert key == "resume.docx"
            monkeypatch.setattr("app.worker.utcnow", lambda: later)
            with ThreadPoolExecutor(max_workers=1) as pool:
                assert pool.submit(claim_job, factory, settings).result(timeout=5) is None
            return {"backend": "supabase", "bucket": settings.supabase_bucket, "object_key": key}

    monkeypatch.setattr("app.storage.get_storage", lambda settings: Storage())
    assert process_one(factory, settings)
    assert get_job(factory).attempts == 1
    with factory() as session:
        assert session.scalar(select(Artifact)) is None
    assert claim_job(factory, settings) is not None
