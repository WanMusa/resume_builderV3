import copy
import json
import shutil
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pydantic import ValidationError

from app.generation import (Letter, MockProvider, OpenAIProvider, ResumeVariables,
                            generate_package, prepare, validate_base)
from app.rendering import PackageFormatError, render_package
from app.storage import DOCX_CONTENT_TYPE, SupabaseStorage, get_storage

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def settings(tmp_path):
    shutil.copyfile(ROOT / "assets/base_resume_example.json", tmp_path / "base_resume.json")
    return SimpleNamespace(assets_dir=tmp_path, prompts_dir=ROOT / "prompts",
                           provider="mock", openai_model="example-model", openai_api_key=None)


@pytest.fixture
def job():
    return json.loads((ROOT / "examples/job-envelope.json").read_text())["jobs"][0]


def filenames(package, job):
    prefix = f'{package["resume"]["identity"]["name"]}-{job["title"]}-{job["company"]}'
    return [f"{prefix}-Resume.docx", f"{prefix}-Cover Letter.docx"]


def test_mock_package_fixed_facts_render_and_storage(settings, job, tmp_path, monkeypatch):
    original = (settings.assets_dir / "base_resume.json").read_bytes()
    package = generate_package(job, settings)
    base = json.loads(original)
    assert package["format_version"] == 4
    assert package["source_resume"] == base
    for key in ("identity", "education", "certifications"):
        assert package["resume"][key] == base["fixed"][key]
    assert package["resume"]["experience"] == base["fixed"]["experience"]
    assert package["resume"]["identity"]["website"] == base["fixed"]["identity"]["website"]
    assert all(role["summary"] and role["highlights"] for role in package["resume"]["experience"])
    assert "professional_experience" not in package["resume"]
    assert (settings.assets_dir / "base_resume.json").read_bytes() == original
    assert package["provenance"]["usage"] == {"resume": {}, "cover_letter": {}}
    files = render_package(package, ROOT / "assets", job)
    assert list(files) == filenames(package, job)
    for content in files.values():
        assert content.startswith(b"PK")
        assert "NOT FOR SUBMISSION" in "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
    assert not (tmp_path / "rendered").exists()
    requests = []

    def post(self, url, **kwargs):
        requests.append((url, kwargs))
        return httpx.Response(200, json={})

    monkeypatch.setattr(httpx.Client, "post", post)
    storage = SupabaseStorage(SimpleNamespace(supabase_url="https://storage.example",
                              supabase_key="test-only", supabase_bucket="private-artifacts"))
    for name, content in files.items():
        result = storage.upload(content, name)
        assert result["object_key"] == name and result["local_path"] is None
        assert requests[-1][1]["content"] == content
        assert requests[-1][1]["headers"]["Content-Type"] == DOCX_CONTENT_TYPE
    assert len(requests) == 2


def test_word_renders_fixed_experience_and_tailored_highlights(settings, job):
    package = json.loads(json.dumps(generate_package(job, settings)))
    resume = package["resume"]
    fixed = package["source_resume"]["fixed"]
    fixed["identity"].update(email="candidate@example.com", mobile="1234",
                              location="Example City", linkedin="https://example.com/profile")
    role = fixed["experience"][0]
    role.update(summary="Fixed role summary.", highlights=["Fixed role contribution."])
    resume["career_profile"] = "Concise tailored profile."
    resume["areas_of_expertise"] = [{"area": "Systems integration",
                                    "description": "Connects business platforms through reliable interfaces.",
                                    "evidence_ids": ["skill-analysis"]}]
    fixed["identity"]["website"] = "https://musakamil.com/"
    fixed["education_records"] = [{"qualification": "Fixed qualification",
                                   "institution": "Fixed institution", "period": "2020",
                                   "details": "Fixed education details"}]
    fixed["education"] = ["DUPLICATE EDUCATION STRING"]
    fixed["references"] = "Fixed references"
    fixed["certifications"] = ["Fixed certification (2021)"]
    fixed["selected_projects_intro"] = (
        "Server-side solutions run on a private Oracle Cloud VPS using Docker; "
        "job-ad capture runs locally in Chrome."
    )
    fixed["selected_projects_url"] = "https://musakamil.com/"
    fixed["selected_projects"] = [
        {"name": "Telegram – Gmail Assistant",
         "description": "An event-driven Gmail assistant with human approval in Telegram."},
        {"name": "Local Web Scraping",
         "description": "A Chrome extension that captures job adverts locally without an LLM."},
    ]
    resume["selected_projects"] = [{"name": "Model copy", "description": "MODEL REWRITTEN AI PROJECT"}]
    resume["identity"]["website"] = "https://model-rewritten.invalid/"
    resume["experience"][0]["summary"] = "MODEL REWRITTEN EXPERIENCE"
    resume["experience"][0]["highlights"] = ["MODEL REWRITTEN ACHIEVEMENT"]
    resume["selected_projects"][0]["description"] = "MODEL REWRITTEN AI PROJECT"
    package["source_resume"]["fixed"]["career_profile"] = "SOURCE PROFILE CONTEXT"
    package["provenance"].update(provider="openai", is_example=False)
    package["review_required"] = True
    files = render_package(package, ROOT / "assets", job)
    doc = Document(BytesIO(files[filenames(package, job)[0]]))
    text = "\n".join(p.text for p in doc.paragraphs)
    for value in ("Fixed role summary.", "Fixed role contribution.", "Concise tailored profile.",
                  "Fixed qualification", "Fixed institution", "Fixed education details",
                  "Fixed references", "Fixed certification", role["employer"], role["dates"],
                  "AGENTIC AI & AUTOMATION SOLUTIONS", fixed["selected_projects_intro"],
                  "Telegram – Gmail Assistant", "An event-driven Gmail assistant with human approval in Telegram.",
                  "Local Web Scraping", "A Chrome extension that captures job adverts locally without an LLM.",
                  "Further project details:"):
        assert value in text
    for value in ("SOURCE ROLE CONTEXT", "SOURCE PLATFORM CONTEXT", "DUPLICATE", "GENERATED DRAFT",
                  "MODEL REWRITTEN EXPERIENCE", "MODEL REWRITTEN ACHIEVEMENT", "MODEL REWRITTEN AI PROJECT",
                  "model-rewritten.invalid"):
        assert value not in text
    assert "PRIVATE-CITATION-ID" not in text
    assert text.count("Fixed qualification") == 1
    name = next(p for p in doc.paragraphs if p.text == resume["identity"]["name"])
    assert name.runs[0].font.name == "Lato" and name.runs[0].font.size.pt == 18
    assert str(name.runs[0].font.color.rgb) == "365F91"
    competency_bullet = next(p for p in doc.paragraphs if "Systems integration:" in p.text)
    experience_bullet = next(p for p in doc.paragraphs if "Fixed role contribution." in p.text)
    key_heading = next(p for p in doc.paragraphs if p.text == "Key Achievements")
    assert competency_bullet.paragraph_format.left_indent == experience_bullet.paragraph_format.left_indent
    assert competency_bullet.paragraph_format.first_line_indent == experience_bullet.paragraph_format.first_line_indent
    project_bullet = next(p for p in doc.paragraphs if "Telegram – Gmail Assistant" in p.text)
    assert project_bullet.text.startswith("\u2022  ") and project_bullet.runs[0].bold
    assert project_bullet.paragraph_format.left_indent == experience_bullet.paragraph_format.left_indent
    assert project_bullet.paragraph_format.first_line_indent == experience_bullet.paragraph_format.first_line_indent
    assert key_heading.runs[0].bold and key_heading.runs[0].italic
    assert text.index("Key Achievements") < text.index("Fixed role contribution.")
    further_link = next(p for p in doc.paragraphs if p.text.startswith("Further project details:"))
    assert further_link.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert further_link.paragraph_format.line_spacing == 1.5
    assert further_link.paragraph_format.space_before.pt == 12
    assert further_link.paragraph_format.space_after.pt == 12
    further_hyperlink = further_link._p.find(qn("w:hyperlink")).find(qn("w:r"))
    further_run_properties = further_hyperlink.find(qn("w:rPr"))
    assert further_run_properties.find(qn("w:i")) is not None
    assert further_run_properties.find(qn("w:sz")).get(qn("w:val")) == "24"
    separators = [p for p in doc.paragraphs if not p.text and p._p.pPr is not None
                  and (border := p._p.pPr.find(qn("w:pBdr"))) is not None
                  and (bottom := border.find(qn("w:bottom"))) is not None
                  and bottom.get(qn("w:color")) == "C9CDD2"]
    assert len(separators) == len(fixed["experience"])
    headings = [p for p in doc.paragraphs if p.text in (
        "CAREER PROFILE", "CORE COMPETENCIES", "AGENTIC AI & AUTOMATION SOLUTIONS",
        "PROFESSIONAL EXPERIENCE",
        "EDUCATION", "REFERENCES")]
    assert len(headings) == 6
    assert all(p.paragraph_format.keep_with_next for p in headings)
    assert text.index("CORE COMPETENCIES") < text.index("AGENTIC AI & AUTOMATION SOLUTIONS")
    assert text.index("AGENTIC AI & AUTOMATION SOLUTIONS") < text.index("PROFESSIONAL EXPERIENCE")
    assert "AREAS OF EXPERTISE/HIGHLIGHTS" not in text
    template = Document(str(ROOT / "assets/templates/resume_template.docx"))
    assert doc.sections[0].left_margin == template.sections[0].left_margin
    assert doc.sections[0].page_width == template.sections[0].page_width
    targets = [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]
    assert "mailto:candidate@example.com" in targets and "https://example.com/profile" in targets
    assert "https://musakamil.com/" in targets
    letter = Document(BytesIO(files[filenames(package, job)[1]]))
    letter_text = "\n".join(p.text for p in letter.paragraphs)
    assert job["company"] in letter_text and "Dear Hiring Manager," in letter_text
    assert all(p in letter_text for p in package["cover_letter"]["paragraphs"])
    assert any(r.font.name == "Aptos" for p in letter.paragraphs for r in p.runs)


def test_filenames_follow_v2_without_directories(settings, job):
    package = generate_package(job, settings)
    package["resume"]["identity"]["name"] = "A/B: Candidate"
    files = render_package(package, ROOT / "assets",
                           {**job, "title": 'Senior: Analyst / AI', "company": 'A|B "Ltd"'})
    assert set(files) == {
        "AB Candidate-Senior Analyst  AI-AB Ltd-Resume.docx",
        "AB Candidate-Senior Analyst  AI-AB Ltd-Cover Letter.docx",
    }
    assert all("/" not in name and "\\" not in name for name in files)


def test_legacy_checkpoint_requires_explicit_regeneration():
    with pytest.raises(PackageFormatError, match="regenerate"):
        render_package({"resume": {"summary": "Old selection-only content"}}, ROOT / "assets")


def test_existing_v2_checkpoint_still_renders(settings, job):
    package = generate_package(job, settings)
    package["format_version"] = 2
    package["resume"]["target_role_tags"] = ["Legacy format"]
    package["resume"]["experience"][0].update(
        tailored_summary="Saved V2 tailored summary.",
        tailored_achievements=[{"highlight": "Saved V2 achievement.",
                                "detail": "Legacy checkpoint content remains renderable.",
                                "evidence_ids": ["legacy-evidence"]}],
    )

    files = render_package(package, ROOT / "assets", job)
    doc = Document(BytesIO(files[filenames(package, job)[0]]))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

    assert "AREAS OF EXPERTISE/HIGHLIGHTS" in text
    assert "Saved V2 tailored summary." in text
    assert "Saved V2 achievement." in text


def test_existing_v3_checkpoint_still_renders_its_generated_achievements(settings, job):
    package = generate_package(job, settings)
    package["format_version"] = 3
    package["resume"]["experience"][0].update(
        tailored_summary="Saved V3 tailored role summary.",
        tailored_achievements=[{"highlight": "Saved V3 experience achievement.",
                                "detail": "Keep legacy role content renderable.",
                                "evidence_ids": ["legacy-role-evidence"]}],
    )
    package["resume"]["key_achievements"] = [{
        "highlight": "Saved V3 achievement.",
        "detail": "Keep this existing checkpoint renderable.",
        "evidence_ids": ["legacy-evidence"],
    }]

    files = render_package(package, ROOT / "assets", job)
    doc = Document(BytesIO(files[filenames(package, job)[0]]))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

    assert "Saved V3 tailored role summary." in text
    assert "Saved V3 experience achievement." in text
    assert "Keep legacy role content renderable." in text
    assert "Saved V3 achievement." in text


def test_v4_renderer_ignores_generated_copies_of_fixed_sections(settings, job):
    package = json.loads(json.dumps(generate_package(job, settings)))
    package["source_resume"]["fixed"]["selected_projects"] = [
        {"name": "Fixed project", "description": "Fixed AI project description."}
    ]
    package["resume"]["selected_projects"] = [
        {"name": "Model copy", "description": "MODEL REWRITTEN PROJECT"}
    ]
    package["resume"]["identity"]["website"] = "https://model-rewritten.invalid/"
    package["resume"]["experience"][0]["summary"] = "MODEL REWRITTEN EXPERIENCE"
    package["resume"]["experience"][0]["highlights"] = ["MODEL REWRITTEN HIGHLIGHT"]
    package["resume"]["selected_projects"][0]["description"] = "MODEL REWRITTEN PROJECT"

    files = render_package(package, ROOT / "assets", job)
    doc = Document(BytesIO(files[filenames(package, job)[0]]))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    links = [relationship.target_ref for relationship in doc.part.rels.values() if relationship.is_external]
    fixed = package["source_resume"]["fixed"]

    assert fixed["identity"]["website"] in links
    assert fixed["experience"][0]["summary"] in text
    assert fixed["experience"][0]["highlights"][0] in text
    assert "Fixed AI project description." in text
    assert "model-rewritten.invalid" not in text
    assert "MODEL REWRITTEN" not in text


def test_model_cannot_modify_fixed_fields(settings, job):
    _, context, prompts = prepare(job, settings)
    output, _ = MockProvider().generate("resume", context, prompts, ResumeVariables)
    with pytest.raises(ValidationError):
        ResumeVariables.model_validate({**output.model_dump(), "identity": {"name": "Invented"}})
    with pytest.raises(ValidationError):
        ResumeVariables.model_validate({**output.model_dump(), "professional_experience": []})
    with pytest.raises(ValidationError):
        ResumeVariables.model_validate({**output.model_dump(), "key_achievements": []})


def test_compact_context_and_fixed_role_order(settings, job):
    path = settings.assets_dir / "base_resume.json"
    base = json.loads(path.read_text())
    base["fixed"]["experience"].append({"id": "role-2", "title": "Engineer",
                                        "employer": "Example", "dates": "2019-2022",
                                        "summary": "Long background context",
                                        "highlights": ["Fixed engineering contribution"]})
    base["fixed"]["career_profile"] = "Original source profile"
    base["fixed"]["selected_projects"] = [{"name": "Project context", "details": "Complete details"}]
    base["achievements"].extend([
        {"id": f"evidence-{i}", "experience_id": "role-2", "text": f"Contribution {i}",
         "supporting_details": ["Detailed factual background"]}
        for i in range(20)
    ])
    path.write_text(json.dumps(base))
    _, context, prompts = prepare({**job, "notes": "private operator note"}, settings)
    assert context["evidence"]["achievements"] == base["achievements"]
    assert context["background"]["experience"] == [
        {key: role[key] for key in ("id", "employer", "title", "dates")}
        for role in base["fixed"]["experience"]
    ]
    assert "selected_projects" not in context["background"]
    assert all("platform_description" not in role and "summary" not in role
               for role in context["background"]["experience"])
    assert "identity" not in context["background"] and "notes" not in context["job"]
    assert "untrusted" in prompts["resume_system"]
    assert "ATS-compatible" in prompts["resume_system"]
    assert "capability-level inference" in prompts["resume_system"]
    assert "hands-on use" in prompts["resume_system"]
    package = generate_package(job, settings)
    assert package["format_version"] == 4
    assert [r["id"] for r in package["resume"]["experience"]] == ["role-1", "role-2"]
    assert package["resume"]["experience"] == base["fixed"]["experience"]
    assert package["source_resume"] == base


@pytest.mark.parametrize("invalid,reason", [
    ("expertise", "unknown_evidence"), ("letter", "unknown_evidence"),
    ("long-profile", "content_length"), ("long-expertise", "content_length"),
    ("long-letter", "content_length"),
])
def test_rejects_invalid_or_overlong_prose(settings, job, monkeypatch, invalid, reason):
    original = MockProvider.generate

    def alter(self, kind, context, prompts, schema):
        result, usage = original(self, kind, context, prompts, schema)
        if kind == "resume":
            if invalid == "expertise":
                result.areas_of_expertise[0].evidence_ids = ["invented"]
            elif invalid == "long-profile":
                result.career_profile = "word " * 151
            elif invalid == "long-expertise":
                result.areas_of_expertise[0].area = "Analysis"
                result.areas_of_expertise[0].description = "word " * 26
        elif invalid == "letter":
            result.evidence_ids = ["invented"]
        elif invalid == "long-letter":
            result.paragraphs = ["word " * 180, "word " * 180]
        return result, usage

    monkeypatch.setattr(MockProvider, "generate", alter)
    with pytest.raises(ValueError) as error:
        generate_package(job, settings)
    assert error.value.reason == reason


def test_base_rejects_duplicate_and_unknown_ids(settings):
    base = json.loads((settings.assets_dir / "base_resume.json").read_text())
    base["achievements"].append(dict(base["achievements"][0]))
    with pytest.raises(ValueError, match="unique"):
        validate_base(base)
    base["achievements"].pop()
    base["achievements"][0]["experience_id"] = "unknown"
    with pytest.raises(ValueError, match="unknown experience"):
        validate_base(base)


@pytest.mark.parametrize("field", ["profile", "expertise", "letter"])
@pytest.mark.parametrize("excess", [0, 1])
def test_word_limit_boundaries(settings, job, monkeypatch, field, excess):
    original = MockProvider.generate

    def boundary(self, kind, context, prompts, schema):
        result, usage = original(self, kind, context, prompts, schema)
        if kind == "resume":
            if field == "profile":
                result.career_profile = "word " * (150 + excess)
            elif field == "expertise":
                result.areas_of_expertise[0].area = "Analysis"
                result.areas_of_expertise[0].description = "word " * (25 + excess)
        elif field == "letter":
            result.paragraphs = ["word " * 175, "word " * (175 + excess)]
        return result, usage

    monkeypatch.setattr(MockProvider, "generate", boundary)
    if excess:
        with pytest.raises(ValueError) as error:
            generate_package(job, settings)
        assert error.value.reason == "content_length"
    else:
        assert generate_package(job, settings)["format_version"] == 4


@pytest.mark.parametrize("excess", [0, 1])
def test_combined_resume_word_budget(settings, job, monkeypatch, excess):
    path = settings.assets_dir / "base_resume.json"
    base = json.loads(path.read_text())
    base["fixed"]["experience"] = [
        {"id": f"role-{i}", "employer": "Example", "title": "Analyst", "dates": "2020-2021",
         "summary": "Fixed source experience summary.",
         "highlights": ["Fixed source experience contribution."]}
        for i in range(5)
    ]
    base["expertise"] = [{"id": f"skill-{i}", "text": "Analysis"} for i in range(7)]
    base["achievements"] = [
        {"id": f"achievement-{i}-{j}", "experience_id": f"role-{i}", "text": "Delivered analysis"}
        for i in range(5) for j in range(2)
    ]
    path.write_text(json.dumps(base))
    original = MockProvider.generate

    def boundary(self, kind, context, prompts, schema):
        if kind != "resume":
            return original(self, kind, context, prompts, schema)
        result = {
            "career_profile": "word " * (150 + excess),
            "areas_of_expertise": [
                {"area": "Analysis", "description": "word " * 25, "evidence_ids": [r["id"]]}
                for r in base["expertise"]
            ],
        }
        return schema.model_validate(result), {}

    monkeypatch.setattr(MockProvider, "generate", boundary)
    if excess:
        with pytest.raises(ValueError) as error:
            generate_package(job, settings)
        assert error.value.reason == "content_length"
    else:
        assert generate_package(job, settings)["format_version"] == 4


def test_model_output_contains_only_profile_and_competencies(settings, job):
    _, context, prompts = prepare(job, settings)
    context["evidence"]["expertise"] = [
        {"id": f"skill-{index}", "text": f"Competency {index}"} for index in range(14)
    ]
    output, _ = MockProvider().generate("resume", context, prompts, ResumeVariables)

    assert set(output.model_dump()) == {"career_profile", "areas_of_expertise"}
    assert len(output.areas_of_expertise) == 7


def test_competencies_cannot_cite_unknown_evidence(settings, job, monkeypatch):
    original = MockProvider.generate

    def unknown_citation(self, kind, context, prompts, schema):
        result, usage = original(self, kind, context, prompts, schema)
        if kind == "resume":
            result.areas_of_expertise[0].evidence_ids = ["unknown-evidence"]
        return result, usage

    monkeypatch.setattr(MockProvider, "generate", unknown_citation)
    with pytest.raises(ValueError) as error:
        generate_package(job, settings)
    assert error.value.reason == "unknown_evidence"


def test_lengthy_source_is_retained_without_expanding_mock_output(settings, job):
    path = settings.assets_dir / "base_resume.json"
    base = json.loads(path.read_text())
    base["achievements"].extend([
        {"id": f"additional-{i}", "experience_id": "role-1", "text": "Delivered factual analysis",
         "supporting_details": ["Complete background context"], "critical": i == 149}
        for i in range(150)
    ])
    path.write_text(json.dumps(base))
    package = generate_package(job, settings)
    assert package["source_resume"] == base
    assert "key_achievements" not in package["resume"]
    assert package["resume"]["experience"] == base["fixed"]["experience"]


@pytest.mark.parametrize("kind,schema", [("resume", ResumeVariables), ("cover_letter", Letter)])
@pytest.mark.parametrize("empty_expertise", [False, True])
def test_openai_schema_and_prose_response(settings, job, monkeypatch, kind, schema, empty_expertise):
    settings.openai_api_key = "test-placeholder"
    settings.openai_base_url = "https://provider.example/v1"
    _, context, prompts = prepare(job, settings)
    if empty_expertise:
        context["evidence"]["expertise"] = []
    calls = []

    def post(self, url, **kwargs):
        assert kwargs["headers"]["Authorization"].split() == ["Bearer", settings.openai_api_key]
        calls.append(kwargs["json"])
        request_context = json.loads(kwargs["json"]["messages"][1]["content"].rsplit("\n", 1)[1])
        result, _ = MockProvider().generate(kind, request_context, prompts, schema)
        return httpx.Response(200, json={"choices": [{"message": {
            "content": result.model_dump_json()}}], "usage": {"prompt_tokens": 100}})

    monkeypatch.setattr(httpx.Client, "post", post)
    original_schema = copy.deepcopy(schema.model_json_schema())
    for suffix in ("first", "second"):
        context["evidence"]["achievements"][0]["id"] = f"evidence-{suffix}"
        output, usage = OpenAIProvider(settings).generate(kind, context, prompts, schema)
        assert usage["prompt_tokens"] == 100
        contract = calls[-1]["response_format"]["json_schema"]
        assert contract["strict"] and not contract["schema"]["additionalProperties"]
        if kind == "resume":
            definitions = contract["schema"]["$defs"]
            assert set(contract["schema"]["properties"]) == {"career_profile", "areas_of_expertise"}
            assert "Achievement" not in definitions
            assert output.career_profile
            assert "professional_experience" not in contract["schema"]["properties"]
            assert "target_role_tags" not in contract["schema"]["properties"]
            if empty_expertise:
                assert contract["schema"]["properties"]["areas_of_expertise"]["maxItems"] == 0
            else:
                assert contract["schema"]["properties"]["areas_of_expertise"]["maxItems"] == 7
                assert definitions["Expertise"]["properties"]["evidence_ids"]["items"]["enum"] == [
                    r["id"] for r in context["evidence"]["expertise"]]
        else:
            assert contract["schema"]["properties"]["evidence_ids"]["items"]["enum"] == [
                f"evidence-{suffix}"]
    assert schema.model_json_schema() == original_schema and len(calls) == 2


def test_final_storage_has_no_local_fallback():
    with pytest.raises(ValueError, match="local output is disabled"):
        get_storage(SimpleNamespace(storage_backend="local"))
    with pytest.raises(ValueError, match="configured"):
        get_storage(SimpleNamespace(storage_backend="supabase", supabase_url=None,
                                    supabase_key=None, supabase_bucket=None))


@pytest.mark.parametrize("public", [False, True, None])
def test_supabase_requires_private_bucket(monkeypatch, public):
    storage = SupabaseStorage(SimpleNamespace(supabase_url="https://storage.example",
                              supabase_key="test-only", supabase_bucket="private-artifacts"))
    monkeypatch.setattr(httpx.Client, "get",
                        lambda *args, **kwargs: httpx.Response(200, json={"public": public}))
    if public is False:
        storage.check_bucket()
    else:
        with pytest.raises(ValueError, match="private"):
            storage.check_bucket()


@pytest.mark.parametrize("key", ["folder/resume.docx", "../escape.docx", r"folder\resume.docx",
                                "resume.md", "bad\x00.docx"])
def test_supabase_rejects_subdirectories_and_unsafe_keys(key):
    storage = SupabaseStorage(SimpleNamespace(supabase_url="https://storage.example",
                              supabase_key="test-only", supabase_bucket="private-artifacts"))
    with pytest.raises(ValueError, match="safe Word"):
        storage.upload(b"document", key)


def test_supabase_signed_download_url_is_short_lived_and_same_origin(monkeypatch):
    settings = SimpleNamespace(supabase_url="https://storage.example",
                               supabase_key="test-only", supabase_bucket="private-artifacts")
    calls = []

    def post(self, url, **kwargs):
        calls.append((url, kwargs))
        return httpx.Response(200, json={"signedURL": "/object/sign/private-artifacts/Candidate-Resume.docx?token=temporary"})

    monkeypatch.setattr(httpx.Client, "post", post)
    monkeypatch.setattr(SupabaseStorage, "check_bucket", lambda self: None)
    storage = SupabaseStorage(settings)

    signed_url = storage.signed_download_url("Candidate-Resume.docx")

    assert signed_url == "https://storage.example/storage/v1/object/sign/private-artifacts/Candidate-Resume.docx?token=temporary"
    assert calls[0][1]["json"] == {"expiresIn": 300}
    assert calls[0][1]["headers"]["apikey"] == "test-only"
    with pytest.raises(ValueError, match="safe root-level"):
        storage.signed_download_url("../private/Resume.docx")


def test_supabase_signed_download_url_rejects_foreign_origins(monkeypatch):
    settings = SimpleNamespace(supabase_url="https://storage.example",
                               supabase_key="test-only", supabase_bucket="private-artifacts")
    monkeypatch.setattr(SupabaseStorage, "check_bucket", lambda self: None)
    monkeypatch.setattr(httpx.Client, "post", lambda *args, **kwargs: httpx.Response(
        200, json={"signedURL": "https://attacker.example/file.docx"},
    ))

    with pytest.raises(RuntimeError, match="unexpected origin"):
        SupabaseStorage(settings).signed_download_url("Candidate-Resume.docx")


def test_supabase_signed_download_url_refuses_public_bucket(monkeypatch):
    settings = SimpleNamespace(supabase_url="https://storage.example",
                               supabase_key="test-only", supabase_bucket="private-artifacts")
    monkeypatch.setattr(httpx.Client, "get", lambda *args, **kwargs: httpx.Response(
        200, json={"public": True},
    ))
    monkeypatch.setattr(httpx.Client, "post", lambda *args, **kwargs: pytest.fail(
        "must not sign an object in a public bucket",
    ))

    with pytest.raises(ValueError, match="private"):
        SupabaseStorage(settings).signed_download_url("Candidate-Resume.docx")


def test_supabase_failure_is_sanitized_and_overwrite_key_stable(monkeypatch):
    settings = SimpleNamespace(supabase_url="https://storage.example", supabase_key="test-only",
                               supabase_bucket="private-artifacts")
    calls = []

    def post(self, url, **kwargs):
        calls.append((url, kwargs))
        return httpx.Response(503, text="sensitive provider response")

    monkeypatch.setattr(httpx.Client, "post", post)
    for _ in range(2):
        with pytest.raises(RuntimeError, match="HTTP 503") as error:
            SupabaseStorage(settings).upload(b"word-document", "Candidate-Role-Company-Resume.docx")
        assert "sensitive" not in str(error.value)
    assert calls[0][0] == calls[1][0]
    assert calls[0][1]["headers"]["x-upsert"] == "true"
