"""In-memory Word rendering using V2 templates and Python-controlled formatting."""

from datetime import date
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Lato"
BLUE = RGBColor(0x36, 0x5F, 0x91)


class PackageFormatError(ValueError):
    def __init__(self):
        super().__init__("Unsupported saved package; regenerate via /jobs/{id}/regenerate")


def _document(template: Path):
    doc = Document(str(template))
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)
    return doc


def _paragraph(doc, text: str = "", *, bold: bool = False, size: int = 10,
               font: str = FONT, after: int = 4, before: int = 0,
               justify: bool = False, italic: bool = False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.widow_control = True
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


def _section(doc, title: str):
    paragraph = _paragraph(doc, title, bold=True, size=12, before=14, after=2)
    paragraph.runs[0].font.color.rgb = BLUE
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for name, value in (("val", "single"), ("sz", "6"), ("space", "0"), ("color", "365F91")):
        bottom.set(qn(f"w:{name}"), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)
    paragraph.paragraph_format.keep_with_next = True


def _bullet(doc, highlight: str, detail: str):
    paragraph = _paragraph(doc, "\u2022  " + highlight, bold=True, justify=True)
    if detail:
        run = paragraph.add_run(" " + detail)
        run.font.name, run.font.size = FONT, Pt(10)
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    paragraph.paragraph_format.keep_together = True


def _hyperlink(paragraph, label: str, url: str, font: str, size: int):
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    for tag, values in (
        ("rFonts", {"ascii": font, "hAnsi": font}),
        ("sz", {"val": str(size * 2)}),
        ("color", {"val": "0563C1"}),
        ("u", {"val": "single"}),
    ):
        element = OxmlElement(f"w:{tag}")
        for name, value in values.items():
            element.set(qn(f"w:{name}"), value)
        properties.append(element)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _identity(doc, identity: dict, *, cover: bool = False, tags=()):
    font, size = ("Aptos", 11) if cover else (FONT, 10)
    if cover:
        _paragraph(doc, identity["name"], bold=True, size=12, font=font, after=0)
        for key in ("location", "email", "mobile", "website"):
            value = identity.get(key)
            if not value:
                continue
            paragraph = _paragraph(doc, font=font, size=size, after=0)
            if key == "email":
                _hyperlink(paragraph, value, f"mailto:{value}", font, size)
            elif key == "website":
                label = value.removeprefix("https://").removeprefix("www.").rstrip("/")
                _hyperlink(paragraph, label, value, font, size)
            else:
                run = paragraph.add_run(value)
                run.font.name, run.font.size = font, Pt(size)
        return

    name = _paragraph(doc, identity["name"], bold=True, size=18, font=font)
    name.runs[0].font.color.rgb = BLUE
    if tags:
        _paragraph(doc, "  |  ".join(tags), bold=True, after=2)
    contact = _paragraph(doc, font=font, size=size, after=6)
    parts = []
    for key in ("email", "mobile", "location", "linkedin", "website"):
        value = identity.get(key)
        if value:
            if parts:
                run = contact.add_run(" | ")
                run.font.name, run.font.size = font, Pt(size)
            if key in ("email", "linkedin", "website"):
                if key == "email":
                    label, url = value, f"mailto:{value}"
                elif key == "linkedin":
                    label, url = "LinkedIn", value
                else:
                    label, url = value.removeprefix("https://").removeprefix("www.").rstrip("/"), value
                _hyperlink(contact, label, url, font, size)
            else:
                run = contact.add_run(value)
                run.font.name, run.font.size = font, Pt(size)
            parts.append(value)
    if not parts and identity.get("contact"):
        run = contact.add_run(identity["contact"])
        run.font.name, run.font.size = font, Pt(size)


def _right_date(doc, paragraph, text: str):
    section = doc.sections[-1]
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        section.page_width - section.left_margin - section.right_margin,
        WD_TAB_ALIGNMENT.RIGHT,
    )
    run = paragraph.add_run("\t" + text)
    run.font.name, run.font.size = FONT, Pt(10)


def _banner(doc, package):
    if package["provenance"]["provider"] == "mock" or package["provenance"]["is_example"]:
        _paragraph(doc, "DEVELOPMENT EXAMPLE - NOT FOR SUBMISSION", bold=True)


def _bytes(doc) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_package(package: dict, assets_dir: Path, job: dict | None = None) -> dict[str, bytes]:
    version = package.get("format_version")
    if version not in (2, 3, 4):
        raise PackageFormatError()
    resume = package["resume"]
    fixed = package.get("source_resume", {}).get("fixed", {}) if version == 4 else resume
    doc = _document(assets_dir / "templates" / "resume_template.docx")
    _banner(doc, package)
    _identity(doc, fixed.get("identity", resume["identity"]),
              tags=resume.get("target_role_tags", ()) if version == 2 else ())
    _section(doc, "CAREER PROFILE")
    _paragraph(doc, resume["career_profile"], justify=True)
    if version == 2:
        _section(doc, "AREAS OF EXPERTISE/HIGHLIGHTS")
        for item in resume["areas_of_expertise"]:
            paragraph = _paragraph(doc, item["area"], bold=True, after=2)
            run = paragraph.add_run(": " + item["description"])
            run.font.name, run.font.size = FONT, Pt(10)
    else:
        _section(doc, "CORE COMPETENCIES")
        for item in resume["areas_of_expertise"]:
            _bullet(doc, f"{item['area']}:", item["description"])
        if version == 3 and resume.get("key_achievements"):
            heading = _paragraph(doc, "Key Achievements", bold=True, before=5, after=1)
            heading.paragraph_format.keep_with_next = True
            for achievement in resume["key_achievements"]:
                _bullet(doc, achievement["highlight"], achievement["detail"])
    projects = fixed.get("selected_projects", []) if version == 4 else resume.get("selected_projects", [])
    if projects:
        _section(doc, "AGENTIC AI & AUTOMATION SOLUTIONS")
        introduction = fixed.get("selected_projects_intro") if version == 4 else resume.get("selected_projects_intro")
        if introduction:
            _paragraph(doc, introduction, justify=True, after=5)
        for project in projects:
            name = project.get("name", "").strip()
            description = project.get("description", "").strip()
            if name and description:
                _bullet(doc, f"{name}:", description)
            elif name or description:
                _bullet(doc, name or description, "")
            for highlight in project.get("highlights", []):
                _bullet(doc, highlight, "")
        project_url = fixed.get("selected_projects_url") if version == 4 else resume.get("selected_projects_url")
        if project_url:
            paragraph = _paragraph(
                doc, "Further project details: ", size=12, before=12, after=12,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.runs[0].italic = True
            _hyperlink(paragraph, "musakamil.com", project_url, FONT, 12)
            link_run = paragraph._p.find(qn("w:hyperlink")).find(qn("w:r"))
            italic = OxmlElement("w:i")
            italic.set(qn("w:val"), "1")
            link_run.find(qn("w:rPr")).append(italic)
    _section(doc, "PROFESSIONAL EXPERIENCE")
    experience = fixed.get("experience", resume["experience"]) if version == 4 else resume["experience"]
    for role in experience:
        paragraph = _paragraph(doc, f"{role['title']}  \u2015  {role['employer']}",
                               bold=True, size=11, before=10, after=2)
        _right_date(doc, paragraph, role["dates"])
        paragraph.paragraph_format.keep_with_next = True
        if version == 2 or (version == 3 and "tailored_summary" in role):
            _paragraph(doc, role["tailored_summary"], justify=True, after=6)
            achievements = role["tailored_achievements"]
            if achievements:
                paragraph = _paragraph(doc, "Key Achievements", bold=True, after=1)
                paragraph.runs[0].italic = True
                paragraph.paragraph_format.keep_with_next = True
            for achievement in achievements:
                _bullet(doc, achievement["highlight"], achievement["detail"])
        else:
            _paragraph(doc, role["summary"], justify=True, after=6)
            highlights = role.get("highlights", [])
            if highlights:
                achievement_heading = _paragraph(doc, "Key Achievements", bold=True, italic=True, after=1)
                achievement_heading.paragraph_format.keep_with_next = True
                for highlight in highlights:
                    _bullet(doc, highlight, "")
        separator = doc.add_paragraph()
        separator.paragraph_format.space_before = Pt(7)
        separator.paragraph_format.space_after = Pt(7)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for name, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "C9CDD2")):
            bottom.set(qn(f"w:{name}"), value)
        border.append(bottom)
        separator._p.get_or_add_pPr().append(border)
    _section(doc, "EDUCATION")
    for education in fixed.get("education_records") or fixed.get("education", []) if version == 4 else resume.get("education_records") or resume.get("education", []):
        if isinstance(education, str):
            _paragraph(doc, education)
        else:
            paragraph = _paragraph(doc, education["qualification"], bold=True, before=3)
            run = paragraph.add_run("  \u2013  " + education["institution"])
            run.font.name, run.font.size = FONT, Pt(10)
            if education.get("period"):
                _right_date(doc, paragraph, education["period"])
            if education.get("details"):
                paragraph.paragraph_format.keep_with_next = True
                details = education["details"]
                _paragraph(doc, details if isinstance(details, str) else "; ".join(details),
                           justify=True)
    for certification in fixed.get("certifications", []) if version == 4 else resume.get("certifications", []):
        year = re.search(r"\s*\((\d{4})\)\s*$", certification)
        paragraph = _paragraph(doc, certification[:year.start()] if year else certification,
                               bold=True)
        if year:
            _right_date(doc, paragraph, year.group(1))
    references = fixed.get("references") if version == 4 else resume.get("references")
    if references:
        _section(doc, "REFERENCES")
        _paragraph(doc, references if isinstance(references, str) else "; ".join(references))

    letter = _document(assets_dir / "templates" / "cover_letter_template.docx")
    _banner(letter, package)
    _identity(letter, resume["identity"], cover=True)
    _paragraph(letter, font="Aptos", size=11, after=0)
    _paragraph(letter, date.today().strftime("%d %B %Y").lstrip("0"),
               font="Aptos", size=11, after=0)
    _paragraph(letter, font="Aptos", size=11, after=0)
    if job and job.get("company"):
        _paragraph(letter, job["company"], font="Aptos", size=11, after=0)
        _paragraph(letter, font="Aptos", size=11, after=0)
        _paragraph(letter, font="Aptos", size=11, after=0)
    _paragraph(letter, "Dear Hiring Manager,", font="Aptos", size=11, after=10)
    for text in package["cover_letter"]["paragraphs"]:
        _paragraph(letter, text, font="Aptos", size=11, after=10, justify=True)
    _paragraph(letter, "Kind regards,", font="Aptos", size=11)
    _paragraph(letter, resume["identity"]["name"], bold=True, font="Aptos", size=11)
    job = job or {}
    parts = [resume["identity"]["name"], job.get("title") or "Role",
             job.get("company") or "Organisation"]
    safe_parts = [re.sub(r'[\\/:*?"<>|\x00-\x1f]', "", part).strip() for part in parts]
    if not all(safe_parts):
        raise ValueError("Candidate, position and organisation must produce nonempty filenames")
    prefix = "-".join(safe_parts)
    return {f"{prefix}-Resume.docx": _bytes(doc), f"{prefix}-Cover Letter.docx": _bytes(letter)}
